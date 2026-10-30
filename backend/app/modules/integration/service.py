import csv
import hashlib
import hmac
import io
import secrets
import uuid
from collections.abc import Sequence
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import set_tenant_context
from app.core.errors import AppError, AuthenticationError, NotFoundError
from app.core.importing import (
    ApplyResult,
    ImportContext,
    ImportSpec,
    ParsedRow,
    RowError,
    validate_row,
)
from app.core.parsing import blank_to_none, normalize_token
from app.core.schemas import PageParams, paginate
from app.modules.integration.models import ApiKey, ImportJob, ImportKind, ImportSource, ImportStatus
from app.modules.integration.schemas import ImportResultOut, RowErrorOut
from app.modules.tenancy.models import Company

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_STORED_ERRORS = 500
DATA_SHEET = "Dữ liệu"

RawRow = tuple[int, dict[str, Any]]  # (số dòng trong file, header gốc → giá trị)


# --- Đọc file ------------------------------------------------------------------------


def read_file(content: bytes, filename: str) -> list[RawRow]:
    if len(content) > MAX_FILE_BYTES:
        raise AppError("File vượt quá 10MB", code="file_too_large")
    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(content)
    elif name.endswith(".csv"):
        rows = _read_csv(content)
    else:
        raise AppError("Chỉ hỗ trợ file .xlsx hoặc .csv", code="unsupported_file")
    if len(rows) > MAX_ROWS:
        raise AppError(f"File vượt quá {MAX_ROWS} dòng", code="too_many_rows")
    return rows


def _read_xlsx(content: bytes) -> list[RawRow]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise AppError("Không đọc được file Excel", code="invalid_file") from exc
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    rows = []
    for row_no, values in enumerate(it, start=2):
        record = {k: v for k, v in zip(keys, values, strict=False) if k}
        if any(blank_to_none(v) is not None for v in record.values()):
            rows.append((row_no, record))
    wb.close()
    return rows


def _read_csv(content: bytes) -> list[RawRow]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AppError("File CSV phải mã hóa UTF-8", code="invalid_file") from exc
    sample = decoded[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(decoded), delimiter=delimiter)
    rows = []
    for row_no, record in enumerate(reader, start=2):
        record = {k.strip(): v for k, v in record.items() if k}
        if any(blank_to_none(v) is not None for v in record.values()):
            rows.append((row_no, record))
    return rows


def map_headers(spec: ImportSpec, raw_rows: list[RawRow]) -> tuple[list[RawRow], list[RowError]]:
    """Đổi header (nhãn hoặc key) sang key của spec; báo thiếu cột bắt buộc."""
    header_map = spec.header_map()
    mapped: list[RawRow] = []
    present: set[str] = set()
    for row_no, record in raw_rows:
        converted = {}
        for header, value in record.items():
            key = header_map.get(normalize_token(str(header)))
            if key is not None:
                converted[key] = value
                present.add(key)
        mapped.append((row_no, converted))
    errors = []
    if raw_rows:
        missing = [c.label for c in spec.columns if c.required and c.key not in present]
        if missing:
            errors.append(RowError(1, None, f"Thiếu cột bắt buộc: {', '.join(missing)}"))
    return mapped, errors


# --- Chạy import ---------------------------------------------------------------------


async def run_import(
    session: AsyncSession,
    spec: ImportSpec,
    raw_rows: list[RawRow],
    *,
    company_id: uuid.UUID,
    source: ImportSource,
    dry_run: bool,
    filename: str | None = None,
    actor_id: uuid.UUID | None = None,
    api_key_id: uuid.UUID | None = None,
) -> ImportResultOut:
    rows, errors = map_headers(spec, raw_rows)
    if not rows:
        errors.append(RowError(0, None, "Không có dòng dữ liệu nào"))

    parsed: list[ParsedRow[Any]] = []
    if not errors:
        for row_no, record in rows:
            data, row_errors = validate_row(spec, row_no, record, spec.row_model)
            errors.extend(row_errors)
            if data is not None:
                parsed.append(ParsedRow(row_no, data))

    result = ApplyResult()
    if not errors:
        result = await _apply(
            session, spec, parsed, ImportContext(company_id, source.value), dry_run
        )
        errors.extend(result.errors)

    status = (
        ImportStatus.FAILED
        if errors
        else (ImportStatus.VALIDATED if dry_run else ImportStatus.SUCCEEDED)
    )
    errors.sort(key=lambda e: e.row)
    job = ImportJob(
        company_id=company_id,
        kind=ImportKind(spec.kind),
        source=source,
        filename=filename,
        dry_run=dry_run,
        status=status,
        total_rows=len(rows),
        created_count=0 if errors else result.created,
        updated_count=0 if errors else result.updated,
        error_count=len(errors),
        errors=[e.__dict__ for e in errors[:MAX_STORED_ERRORS]],
        created_by=actor_id,
        api_key_id=api_key_id,
    )
    session.add(job)
    await session.flush()
    return ImportResultOut(
        job_id=job.id,
        kind=job.kind,
        status=status,
        dry_run=dry_run,
        total_rows=len(rows),
        created=job.created_count,
        updated=job.updated_count,
        errors=[RowErrorOut(**e.__dict__) for e in errors[:MAX_STORED_ERRORS]],
    )


async def _apply(
    session: AsyncSession,
    spec: ImportSpec,
    parsed: Sequence[ParsedRow[Any]],
    ctx: ImportContext,
    dry_run: bool,
) -> ApplyResult:
    """Ghi trong savepoint; rollback nếu dry-run hoặc có lỗi (all-or-nothing)."""
    savepoint = await session.begin_nested()
    try:
        result = await spec.apply(session, ctx, parsed)
        if not result.errors:
            await session.flush()
    except (AppError, IntegrityError) as exc:
        await savepoint.rollback()
        message = exc.message if isinstance(exc, AppError) else "Dữ liệu vi phạm ràng buộc"
        return ApplyResult(errors=[RowError(0, None, message)])
    if result.errors or dry_run:
        await savepoint.rollback()
    else:
        await savepoint.commit()
    return result


async def list_jobs(session: AsyncSession, params: PageParams) -> tuple[list[ImportJob], int]:
    stmt = select(ImportJob).order_by(ImportJob.created_at.desc())
    return await paginate(session, stmt, params)


# --- Template Excel --------------------------------------------------------------------


def build_template(spec: ImportSpec) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = DATA_SHEET
    header_font = Font(bold=True, color="FFFFFF")
    required_fill = PatternFill("solid", fgColor="C62828")
    optional_fill = PatternFill("solid", fgColor="1565C0")
    for col_idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{col.label} *" if col.required else col.label)
        cell.font = header_font
        cell.fill = required_fill if col.required else optional_fill
        ws.cell(row=2, column=col_idx, value=col.example)
        ws.column_dimensions[cell.column_letter].width = max(14, len(col.label) + 6)
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Hướng dẫn")
    guide.append([spec.title])
    guide["A1"].font = Font(bold=True, size=14)
    guide.append([spec.description])
    guide.append(["Cột có dấu * là bắt buộc. Xóa dòng ví dụ trước khi nhập dữ liệu thật."])
    guide.append([])
    guide.append(["Cột", "Mã cột (API)", "Bắt buộc", "Mô tả", "Ví dụ"])
    for cell in guide[5]:
        cell.font = Font(bold=True)
    for col in spec.columns:
        guide.append(
            [col.label, col.key, "Có" if col.required else "", col.description, col.example]
        )
    for letter, width in zip("ABCDE", (24, 20, 10, 60, 24), strict=True):
        guide.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- API key ---------------------------------------------------------------------------

_KEY_PREFIX = "tms"


def _hash_key(key: str) -> str:
    return hashlib.sha256(key.encode()).hexdigest()


async def create_api_key(
    session: AsyncSession, company_id: uuid.UUID, name: str, actor_id: uuid.UUID
) -> tuple[ApiKey, str]:
    prefix = secrets.token_hex(4)
    key = f"{_KEY_PREFIX}_{prefix}_{secrets.token_urlsafe(32)}"
    api_key = ApiKey(
        company_id=company_id,
        name=name,
        key_prefix=prefix,
        key_hash=_hash_key(key),
        created_by=actor_id,
    )
    session.add(api_key)
    await session.flush()
    return api_key, key


async def list_api_keys(session: AsyncSession) -> list[ApiKey]:
    return list(await session.scalars(select(ApiKey).order_by(ApiKey.created_at.desc())))


async def revoke_api_key(session: AsyncSession, key_id: uuid.UUID) -> ApiKey:
    api_key = await session.get(ApiKey, key_id)
    if api_key is None:
        raise NotFoundError("Không tìm thấy API key")
    api_key.is_active = False
    await session.flush()
    return api_key


async def authenticate_api_key(session: AsyncSession, key: str) -> ApiKey:
    """Xác thực key và chuyển session sang công ty của key."""
    parts = key.split("_", 2)
    if len(parts) != 3 or parts[0] != _KEY_PREFIX:
        raise AuthenticationError("API key không hợp lệ")
    # Policy api_key_lookup cho phép đọc đúng key có prefix này trước khi biết công ty.
    await session.execute(
        text("SELECT set_config('app.api_key_prefix', :prefix, true)"), {"prefix": parts[1]}
    )
    api_key = await session.scalar(select(ApiKey).where(ApiKey.key_prefix == parts[1]))
    if (
        api_key is None
        or not api_key.is_active
        or not hmac.compare_digest(api_key.key_hash, _hash_key(key))
    ):
        raise AuthenticationError("API key không hợp lệ hoặc đã bị thu hồi")
    await set_tenant_context(session, company_id=api_key.company_id, user_id=None)
    company = await session.get(Company, api_key.company_id)
    if company is None or not company.is_active:
        raise AuthenticationError("Công ty đã bị đóng")
    api_key.last_used_at = datetime.now(UTC)
    return api_key
