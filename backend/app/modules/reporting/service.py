"""Báo cáo vận hành & chi phí (chỉ đọc) theo khoảng ngày giao.

Tính từ các chuyến đã xuất phát / kết thúc trong khoảng: tỷ lệ giao thành công, giao lại, km,
% tải trọng, chi phí/đơn, chi phí/kg; chia theo xe, ngày, khách hàng.
"""

import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.modules.costing import service as costing
from app.modules.costing.models import TripCost
from app.modules.dispatch.models import ItemStatus, Trip, TripItem, TripOrderStatus
from app.modules.reporting.schemas import (
    CustomerRow,
    DayRow,
    ReasonRow,
    ReportOut,
    Totals,
    VehicleRow,
)

REPORTED = (TripOrderStatus.DISPATCHED, TripOrderStatus.CLOSED)
MAX_DAYS = 366


@dataclass
class _Acc:
    """Cộng dồn số liệu cho 1 nhóm (toàn bộ / 1 xe / 1 ngày)."""

    trips: int = 0
    trips_closed: int = 0
    trips_without_cost: int = 0
    success: int = 0
    failed: int = 0
    cancelled: int = 0
    redeliveries: int = 0
    orders: set[uuid.UUID] = field(default_factory=set)
    delivered_kg: Decimal = Decimal(0)
    loaded_kg: Decimal = Decimal(0)
    capacity_kg: Decimal = Decimal(0)
    distance_km: Decimal = Decimal(0)
    cost: Decimal = Decimal(0)
    # Chỉ tính trên chuyến đã ghi nhận chi phí (chi phí/đơn, chi phí/kg).
    costed_orders: set[uuid.UUID] = field(default_factory=set)
    costed_kg: Decimal = Decimal(0)

    def totals(self) -> Totals:
        done = self.success + self.failed
        kg, cost = float(self.delivered_kg), float(self.cost)
        return Totals(
            trips=self.trips,
            trips_closed=self.trips_closed,
            trips_without_cost=self.trips_without_cost,
            items_success=self.success,
            items_failed=self.failed,
            items_cancelled=self.cancelled,
            success_rate=self.success / done if done else None,
            redeliveries=self.redeliveries,
            orders=len(self.orders),
            delivered_kg=kg,
            distance_km=float(self.distance_km),
            cost=cost,
            load_utilization=float(self.loaded_kg / self.capacity_kg) if self.capacity_kg else None,
            cost_per_order=cost / len(self.costed_orders) if self.costed_orders else None,
            cost_per_kg=cost / float(self.costed_kg) if self.costed_kg else None,
        )


@dataclass
class _Customer:
    code: str
    name: str
    orders: set[uuid.UUID] = field(default_factory=set)
    success: int = 0
    failed: int = 0
    delivered_kg: Decimal = Decimal(0)
    cost: Decimal = Decimal(0)


@dataclass
class ReportData:
    report: ReportOut
    details: list[dict[str, Any]]  # từng mục hàng (cho file Excel)


def _check_range(date_from: date, date_to: date) -> None:
    if date_to < date_from:
        raise AppError("Ngày kết thúc phải sau ngày bắt đầu", code="invalid_range")
    if (date_to - date_from).days >= MAX_DAYS:
        raise AppError(f"Khoảng báo cáo tối đa {MAX_DAYS} ngày", code="invalid_range")


async def build(
    session: AsyncSession,
    date_from: date,
    date_to: date,
    vehicle_id: uuid.UUID | None = None,
    warehouse_id: uuid.UUID | None = None,
) -> ReportData:
    _check_range(date_from, date_to)
    stmt = (
        select(Trip)
        .where(
            Trip.plan_date >= date_from,
            Trip.plan_date <= date_to,
            Trip.order_status.in_(REPORTED),
        )
        .order_by(Trip.plan_date, Trip.code)
    )
    if vehicle_id:
        stmt = stmt.where(Trip.vehicle_id == vehicle_id)
    if warehouse_id:
        stmt = stmt.where(Trip.warehouse_id == warehouse_id)
    trips = list(await session.scalars(stmt))
    trip_ids = [t.id for t in trips]
    costs = {
        c.trip_id: c
        for c in await session.scalars(select(TripCost).where(TripCost.trip_id.in_(trip_ids)))
    }
    allocated = {a.trip_item_id: a.amount for c in costs.values() for a in c.allocations}

    # Lượt giao lại: phần đơn đã có item giao thất bại trước item này.
    part_ids = {i.order_part_id for t in trips for s in t.stops for i in s.items if i.order_part_id}
    failures: dict[uuid.UUID, list[datetime]] = defaultdict(list)
    for part_id, created in await session.execute(
        select(TripItem.order_part_id, TripItem.created_at).where(
            TripItem.order_part_id.in_(part_ids), TripItem.status == ItemStatus.FAILED
        )
    ):
        if part_id is not None:
            failures[part_id].append(created)

    total = _Acc()
    by_vehicle: dict[uuid.UUID, _Acc] = defaultdict(_Acc)
    by_day: dict[date, _Acc] = defaultdict(_Acc)
    customers: dict[uuid.UUID, _Customer] = {}
    reasons: Counter[str] = Counter()
    details: list[dict[str, Any]] = []
    vehicle_names: dict[uuid.UUID, tuple[str, str]] = {}

    for trip in trips:
        vehicle_names[trip.vehicle_id] = (trip.vehicle.code, trip.vehicle.name)
        groups = (total, by_vehicle[trip.vehicle_id], by_day[trip.plan_date])
        cost = costs.get(trip.id)
        if cost is not None:
            distance = cost.distance_km
        elif trip.planned_distance_km is not None:
            distance = trip.planned_distance_km
        else:
            distance = await costing.estimated_distance_km(trip)
        for acc in groups:
            acc.trips += 1
            acc.trips_closed += trip.order_status == TripOrderStatus.CLOSED
            acc.trips_without_cost += cost is None
            acc.loaded_kg += trip.total_weight_kg
            acc.capacity_kg += trip.vehicle.capacity_kg
            acc.distance_km += distance
            acc.cost += cost.total_cost if cost else Decimal(0)
        for stop in trip.stops:
            for item in stop.items:
                status = ItemStatus(item.status)
                part = item.order_part
                redelivery = part is not None and any(
                    f < item.created_at for f in failures.get(part.id, [])
                )
                for acc in groups:
                    acc.success += status == ItemStatus.SUCCESS
                    acc.failed += status == ItemStatus.FAILED
                    acc.cancelled += status == ItemStatus.CANCELLED
                    acc.redeliveries += redelivery and status in costing.CHARGED
                    if status == ItemStatus.SUCCESS:
                        acc.delivered_kg += item.weight_kg
                        if cost is not None:
                            acc.costed_kg += item.weight_kg
                    if part is not None and status in costing.CHARGED:
                        acc.orders.add(part.order_id)
                        if cost is not None:
                            acc.costed_orders.add(part.order_id)
                if status == ItemStatus.FAILED:
                    reasons[(item.reason or "Không ghi lý do").strip()] += 1
                if part is None:
                    continue
                order = part.order
                amount = allocated.get(item.id, Decimal(0))
                if status in costing.CHARGED:
                    c = customers.setdefault(
                        order.customer_id, _Customer(order.customer.code, order.customer.name)
                    )
                    c.orders.add(order.id)
                    c.success += status == ItemStatus.SUCCESS
                    c.failed += status == ItemStatus.FAILED
                    c.cost += amount
                    if status == ItemStatus.SUCCESS:
                        c.delivered_kg += item.weight_kg
                details.append(
                    {
                        "plan_date": trip.plan_date,
                        "trip": trip.code,
                        "vehicle": trip.vehicle.code,
                        "order": order.external_ref,
                        "customer": f"{order.customer.code} - {order.customer.name}",
                        "temp_group": part.temp_group.value,
                        "status": status.value,
                        "reason": item.reason or "",
                        "weight_kg": float(item.weight_kg),
                        "value": float(item.value),
                        "cost": float(amount),
                        "redelivery": redelivery,
                    }
                )

    report = ReportOut(
        date_from=date_from,
        date_to=date_to,
        totals=total.totals(),
        by_vehicle=sorted(
            (
                VehicleRow(
                    vehicle_id=vid,
                    vehicle_code=vehicle_names[vid][0],
                    vehicle_name=vehicle_names[vid][1],
                    **acc.totals().model_dump(),
                )
                for vid, acc in by_vehicle.items()
            ),
            key=lambda r: r.vehicle_code,
        ),
        by_day=[
            DayRow(plan_date=d, **acc.totals().model_dump()) for d, acc in sorted(by_day.items())
        ],
        by_customer=sorted(
            (
                CustomerRow(
                    customer_id=cid,
                    customer_code=c.code,
                    customer_name=c.name,
                    orders=len(c.orders),
                    items_success=c.success,
                    items_failed=c.failed,
                    delivered_kg=float(c.delivered_kg),
                    cost=float(c.cost),
                    cost_per_kg=float(c.cost / c.delivered_kg) if c.delivered_kg else None,
                )
                for cid, c in customers.items()
            ),
            key=lambda r: -r.cost,
        ),
        failure_reasons=[ReasonRow(reason=r, count=n) for r, n in reasons.most_common()],
    )
    return ReportData(report, details)


# --- Xuất Excel ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
TEMP = {"DRY": "Khô", "COLD": "Lạnh/Đông"}
STATUS = {
    "SUCCESS": "Thành công",
    "FAILED": "Thất bại",
    "CANCELLED": "Đã hủy",
    "PENDING": "Chờ giao",
}

_TOTAL_COLUMNS: list[tuple[str, str]] = [
    ("trips", "Số chuyến"),
    ("items_success", "Giao thành công"),
    ("items_failed", "Giao thất bại"),
    ("success_rate", "Tỷ lệ thành công"),
    ("redeliveries", "Lượt giao lại"),
    ("orders", "Số đơn"),
    ("delivered_kg", "Khối lượng giao (kg)"),
    ("distance_km", "Quãng đường (km)"),
    ("load_utilization", "% tải trọng"),
    ("cost", "Chi phí (đ)"),
    ("cost_per_order", "Chi phí/đơn (đ)"),
    ("cost_per_kg", "Chi phí/kg (đ)"),
]
_PERCENT = {"success_rate", "load_utilization"}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(header) + 2)
    ws.freeze_panes = "A2"


def _total_values(t: Totals) -> list[Any]:
    values = []
    for key, _ in _TOTAL_COLUMNS:
        v = getattr(t, key)
        values.append(round(v, 4) if key in _PERCENT and v is not None else v)
    return values


def to_excel(data: ReportData) -> bytes:
    r = data.report
    wb = Workbook()
    assert wb.active is not None
    wb.remove(wb.active)
    labels = [label for _, label in _TOTAL_COLUMNS]
    _sheet(
        wb,
        "Tổng quan",
        ["Từ ngày", "Đến ngày", *labels],
        [[r.date_from, r.date_to, *_total_values(r.totals)]],
    )
    _sheet(
        wb,
        "Theo xe",
        ["Mã xe", "Tên xe", *labels],
        [[v.vehicle_code, v.vehicle_name, *_total_values(v)] for v in r.by_vehicle],
    )
    _sheet(wb, "Theo ngày", ["Ngày", *labels], [[d.plan_date, *_total_values(d)] for d in r.by_day])
    _sheet(
        wb,
        "Theo khách hàng",
        [
            "Mã KH",
            "Tên KH",
            "Số đơn",
            "Thành công",
            "Thất bại",
            "Khối lượng (kg)",
            "Chi phí (đ)",
            "Chi phí/kg (đ)",
        ],
        [
            [
                c.customer_code,
                c.customer_name,
                c.orders,
                c.items_success,
                c.items_failed,
                c.delivered_kg,
                c.cost,
                c.cost_per_kg,
            ]
            for c in r.by_customer
        ],
    )
    _sheet(
        wb,
        "Lý do thất bại",
        ["Lý do", "Số lượt"],
        [[x.reason, x.count] for x in r.failure_reasons],
    )
    _sheet(
        wb,
        "Chi tiết đơn",
        [
            "Ngày",
            "Lệnh",
            "Xe",
            "Số đơn",
            "Khách hàng",
            "Nhóm hàng",
            "Kết quả",
            "Lý do",
            "Khối lượng (kg)",
            "Giá trị (đ)",
            "Chi phí phân bổ (đ)",
            "Giao lại",
        ],
        [
            [
                d["plan_date"],
                d["trip"],
                d["vehicle"],
                d["order"],
                d["customer"],
                TEMP.get(d["temp_group"], d["temp_group"]),
                STATUS.get(d["status"], d["status"]),
                d["reason"],
                d["weight_kg"],
                d["value"],
                d["cost"],
                "Có" if d["redelivery"] else "",
            ]
            for d in data.details
        ],
    )
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            for i, cell in enumerate(row):
                header = headers[i]
                if header in ("Tỷ lệ thành công", "% tải trọng"):
                    cell.number_format = "0.0%"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.##"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
