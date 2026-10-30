import io
from decimal import Decimal
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from app.modules.orders.models import TempGroup
from app.modules.orders.service import LineDraft, split_parts
from tests.conftest import auth


def xlsx(headers: list[str], rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def upload(
    client: AsyncClient, token: str, kind: str, content: bytes, name: str, dry_run: bool = False
) -> dict[str, Any]:
    resp = await client.post(
        f"/api/v1/imports/{kind}?dry_run={str(dry_run).lower()}",
        headers=auth(token),
        files={"file": (name, content)},
    )
    assert resp.status_code == 200, resp.text
    body: dict[str, Any] = resp.json()
    return body


async def seed_master_data(client: AsyncClient, token: str) -> None:
    """Kho, mặt hàng, khách hàng qua file import (CSV + Excel, nhãn tiếng Việt)."""
    warehouses = (
        "Mã kho,Tên kho,Địa chỉ,Vĩ độ,Kinh độ\n"
        "KHO01,Kho Hà Đông,Hà Đông,20.971,105.776\n"
        "KHO02,Kho Long Biên,Long Biên,21.047,105.889\n"
    ).encode()
    assert (await upload(client, token, "warehouses", warehouses, "kho.csv"))["created"] == 2

    products = xlsx(
        ["Mã hàng *", "Tên hàng *", "Loại hàng *", "Đơn vị tính *", "Trọng lượng (kg) *"],
        [
            ["SUA-TT", "Sữa tươi", "Lạnh", "Thùng", 9.2],
            ["KEM-01", "Kem que", "Đông", "Thùng", 6],
            ["NUOC-01", "Nước suối", "khô", "Thùng", "12,5"],
        ],
    )
    assert (await upload(client, token, "products", products, "hang.xlsx"))["created"] == 3

    customers = xlsx(
        [
            "ma_kh",
            "ten_kh",
            "yeu_cau_xe",
            "ngay_giao",
            "ma_diem_giao",
            "dia_chi_diem_giao",
            "vi_do",
            "kinh_do",
            "diem_mac_dinh",
        ],
        [
            ["KH001", "Đại lý A", "Xe máy", "T2, T4, T6", "DG1", "Số 1", 20.97, 105.78, ""],
            ["KH001", "Đại lý A", "Xe máy", "T2, T4, T6", "DG2", "Số 2", 20.96, 105.77, "C"],
            ["KH002", "Đại lý B", "", "", "DG1", "Số 3", 21.0, 105.8, ""],
        ],
    )
    result = await upload(client, token, "customers", customers, "kh.xlsx")
    assert result["status"] == "SUCCEEDED", result
    assert result["created"] == 2


async def test_template_download_roundtrip(client: AsyncClient, admin_token: str) -> None:
    resp = await client.get("/api/v1/imports/templates/customers", headers=auth(admin_token))
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Dữ liệu", "Hướng dẫn"]
    header = [c.value for c in wb["Dữ liệu"][1]]
    assert header[0] == "Mã KH *"

    specs = (await client.get("/api/v1/imports/specs", headers=auth(admin_token))).json()
    assert specs[0]["kind"] == "warehouses"


async def test_template_example_row_imports_cleanly(client: AsyncClient, admin_token: str) -> None:
    """Dòng ví dụ trong template kho phải import được (template tự nhất quán)."""
    resp = await client.get("/api/v1/imports/templates/warehouses", headers=auth(admin_token))
    result = await upload(client, admin_token, "warehouses", resp.content, "t.xlsx", True)
    assert result["status"] == "VALIDATED", result


async def test_customers_import_and_update(client: AsyncClient, admin_token: str) -> None:
    await seed_master_data(client, admin_token)
    customers = (await client.get("/api/v1/customers", headers=auth(admin_token))).json()
    kh1 = next(c for c in customers["items"] if c["code"] == "KH001")
    assert kh1["required_vehicle_type"] == "MOTORBIKE"
    assert kh1["delivery_weekdays"] == [1, 3, 5]
    assert [p["code"] for p in kh1["delivery_points"] if p["is_default"]] == ["DG2"]

    # Import lại: cập nhật, không tạo mới.
    again = (
        "ma_kh,ten_kh,ma_diem_giao,dia_chi_diem_giao,vi_do,kinh_do\n"
        "KH002,Đại lý B mới,DG1,Số 3,21,105.8\n"
    )
    result = await upload(client, admin_token, "customers", again.encode(), "kh.csv")
    assert (result["created"], result["updated"]) == (0, 1)


async def test_validation_errors_are_all_or_nothing(client: AsyncClient, admin_token: str) -> None:
    content = xlsx(
        ["Mã hàng", "Tên hàng", "Loại hàng", "Đơn vị tính", "Trọng lượng (kg)"],
        [
            ["A1", "Hàng A", "Khô", "Thùng", 5],
            ["A2", "", "Nóng", "Thùng", "abc"],
            ["A1", "Hàng A trùng", "Khô", "Thùng", 5],
        ],
    )
    result = await upload(client, admin_token, "products", content, "hang.xlsx")
    assert result["status"] == "FAILED"
    messages = {(e["row"], e["column"]): e["message"] for e in result["errors"]}
    assert messages[(3, "Tên hàng")] == "Bắt buộc nhập"
    assert "Nóng" in messages[(3, "Loại hàng")]
    assert "không phải là số" in messages[(3, "Trọng lượng (kg)")]

    # Dòng 2 hợp lệ nhưng không được ghi vì file có lỗi.
    page = (await client.get("/api/v1/products", headers=auth(admin_token))).json()
    assert page["total"] == 0

    missing_col = await upload(client, admin_token, "products", b"ma_hang\nX\n", "x.csv")
    assert "Thiếu cột bắt buộc" in missing_col["errors"][0]["message"]


async def test_dry_run_does_not_write(client: AsyncClient, admin_token: str) -> None:
    content = "ma_kho;ten_kho;dia_chi;vi_do;kinh_do\nKHO09;Kho thử;HN;21;105\n".encode()
    result = await upload(client, admin_token, "warehouses", content, "k.csv", dry_run=True)
    assert (result["status"], result["created"]) == ("VALIDATED", 1)
    page = (await client.get("/api/v1/warehouses", headers=auth(admin_token))).json()
    assert page["total"] == 0

    jobs = (await client.get("/api/v1/imports/jobs", headers=auth(admin_token))).json()
    assert jobs["items"][0]["status"] == "VALIDATED"


async def test_vehicles_import(client: AsyncClient, admin_token: str) -> None:
    await seed_master_data(client, admin_token)
    content = xlsx(
        [
            "ma_xe",
            "ten_xe",
            "bien_so",
            "loai_xe",
            "trong_tai_kg",
            "chi_phi_km",
            "can_phu_xe",
            "loai_hang_uu_tien",
            "kho_lay_hang",
            "pp_chi_phi",
        ],
        [
            ["XE01", "Xe tải 2 tấn", "29C-001", "Ô tô", 2000, 8000, "C", "", "KHO01,KHO02", ""],
            [
                "XE02",
                "Xe máy 1",
                "29X1-002",
                "Xe máy",
                250,
                2000,
                "K",
                "Hàng khô",
                "KHO01",
                "Chuyến",
            ],
            ["XE03", "Xe lỗi", "29C-003", "Ô tô", 900, 5000, "", "", "KHO99", ""],
        ],
    )
    result = await upload(client, admin_token, "vehicles", content, "xe.xlsx")
    assert result["status"] == "FAILED"
    assert result["errors"] == [
        {"row": 4, "column": "Kho lấy hàng", "message": "Kho không tồn tại: KHO99"}
    ]

    ok = await upload(
        client,
        admin_token,
        "vehicles",
        xlsx(
            [
                "ma_xe",
                "ten_xe",
                "bien_so",
                "loai_xe",
                "trong_tai_kg",
                "chi_phi_km",
                "kho_lay_hang",
                "pp_chi_phi",
                "loai_hang_uu_tien",
                "can_phu_xe",
            ],
            [
                [
                    "XE02",
                    "Xe máy 1",
                    "29X1-002",
                    "Xe máy",
                    250,
                    2000,
                    "KHO01",
                    "Chuyến",
                    "Hàng khô",
                    "K",
                ]
            ],
        ),
        "xe.xlsx",
    )
    assert ok["created"] == 1
    vehicle = (await client.get("/api/v1/vehicles", headers=auth(admin_token))).json()["items"][0]
    assert (vehicle["vehicle_type"], vehicle["cargo_class"], vehicle["cost_allocation_method"]) == (
        "MOTORBIKE",
        "DRY_ONLY",
        "PER_TRIP",
    )


ORDER_HEADERS = [
    "so_don",
    "ngay_dat",
    "ngay_giao_yc",
    "ma_kh",
    "ma_diem_giao",
    "gia_tri_don",
    "ma_hang",
    "so_luong",
    "thanh_tien",
]


async def test_orders_import_splits_by_temperature(client: AsyncClient, admin_token: str) -> None:
    await seed_master_data(client, admin_token)
    content = xlsx(
        ORDER_HEADERS,
        [
            ["DH001", "29/10/2026", "31/10/2026", "KH001", "", 1_000_000, "SUA-TT", 10, None],
            ["DH001", "29/10/2026", "31/10/2026", "KH001", "", 1_000_000, "KEM-01", 5, None],
            ["DH001", "29/10/2026", "31/10/2026", "KH001", "", 1_000_000, "NUOC-01", 4, None],
            ["DH002", "2026-10-29", "2026-10-30", "kh002", "DG1", 500_000, "NUOC-01", 2, None],
        ],
    )
    result = await upload(client, admin_token, "orders", content, "don.xlsx")
    assert result["status"] == "SUCCEEDED", result
    assert result["created"] == 2

    orders = (await client.get("/api/v1/orders?q=DH001", headers=auth(admin_token))).json()["items"]
    parts = {p["temp_group"]: p for p in orders[0]["parts"]}
    # Lạnh 10*9.2 + Đông 5*6 = 122 kg; Khô 4*12.5 = 50 kg → giá trị chia theo khối lượng.
    assert Decimal(parts["COLD"]["weight_kg"]) == Decimal("122")
    assert Decimal(parts["DRY"]["weight_kg"]) == Decimal("50")
    assert Decimal(parts["COLD"]["value"]) + Decimal(parts["DRY"]["value"]) == 1_000_000
    assert all(p["status"] == "PENDING" for p in parts.values())

    # Không ghi điểm giao → dùng điểm mặc định của khách (DG2).
    detail = (
        await client.get(f"/api/v1/orders/{orders[0]['id']}", headers=auth(admin_token))
    ).json()
    customers = (await client.get("/api/v1/customers?q=KH001", headers=auth(admin_token))).json()
    dg2 = next(p for p in customers["items"][0]["delivery_points"] if p["code"] == "DG2")
    assert detail["delivery_point_id"] == dg2["id"]
    assert len(detail["lines"]) == 3

    # Import lại đơn chưa lập kế hoạch → thay thế toàn bộ.
    update = xlsx(
        ORDER_HEADERS,
        [
            ["DH001", "29/10/2026", "31/10/2026", "KH001", "", 300_000, "NUOC-01", 6, 300_000],
        ],
    )
    result = await upload(client, admin_token, "orders", update, "don.xlsx")
    assert (result["created"], result["updated"]) == (0, 1), result["errors"]
    detail = (
        await client.get(f"/api/v1/orders/{orders[0]['id']}", headers=auth(admin_token))
    ).json()
    assert [p["temp_group"] for p in detail["parts"]] == ["DRY"]

    # Hủy đơn rồi import lại → bị khóa.
    resp = await client.post(f"/api/v1/orders/{orders[0]['id']}/cancel", headers=auth(admin_token))
    assert resp.json()["status"] == "CANCELLED"
    result = await upload(client, admin_token, "orders", update, "don.xlsx")
    assert "không thể cập nhật" in result["errors"][0]["message"]


async def test_orders_import_reference_errors(client: AsyncClient, admin_token: str) -> None:
    await seed_master_data(client, admin_token)
    content = xlsx(
        ORDER_HEADERS,
        [
            ["DH010", "29/10/2026", "28/10/2026", "KH999", "", 1000, "SUA-TT", 1, None],
            ["DH011", "29/10/2026", "31/10/2026", "KH001", "DG9", 1000, "XXX", 1, None],
            ["DH012", "29/10/2026", "31/10/2026", "KH001", "", 1000, "SUA-TT", 1, None],
            ["DH012", "29/10/2026", "01/11/2026", "KH001", "", 1000, "KEM-01", 1, None],
        ],
    )
    result = await upload(client, admin_token, "orders", content, "don.xlsx")
    errors = {(e["row"], e["column"]) for e in result["errors"]}
    assert errors == {
        (2, "Mã KH"),
        (2, "Ngày giao yêu cầu"),
        (3, "Mã điểm giao"),
        (3, "Mã hàng"),
        (5, "Ngày giao yêu cầu"),
    }


async def test_inventory_snapshot(client: AsyncClient, admin_token: str) -> None:
    await seed_master_data(client, admin_token)
    first = b"ma_kho,ma_hang,so_luong_ton\nKHO01,SUA-TT,100\nKHO01,KEM-01,50\nKHO02,SUA-TT,7\n"
    assert (await upload(client, admin_token, "inventory", first, "ton.csv"))["created"] == 3
    # Snapshot mới cho KHO01 → KEM-01 biến mất khỏi KHO01, KHO02 giữ nguyên.
    second = b"ma_kho,ma_hang,so_luong_ton\nKHO01,SUA-TT,80\n"
    assert (await upload(client, admin_token, "inventory", second, "ton.csv"))["status"] == (
        "SUCCEEDED"
    )
    from sqlalchemy import select

    from app.core.db import SessionFactory, set_tenant_context
    from app.modules.inventory.models import InventoryBalance

    me = (await client.get("/api/v1/auth/me", headers=auth(admin_token))).json()
    async with SessionFactory() as session:
        import uuid

        await set_tenant_context(session, company_id=uuid.UUID(me["company"]["id"]), user_id=None)
        balances = (await session.scalars(select(InventoryBalance))).all()
    assert sorted(float(b.qty_on_hand) for b in balances) == [7, 80]


async def test_integration_api_with_api_key(client: AsyncClient, admin_token: str) -> None:
    resp = await client.post(
        "/api/v1/api-keys", headers=auth(admin_token), json={"name": "Hệ thống bán hàng"}
    )
    assert resp.status_code == 201
    key = resp.json()["api_key"]

    rows = [
        {
            "ma_hang": "API-1",
            "ten_hang": "Hàng API",
            "loai_hang": "DRY",
            "don_vi_tinh": "Thùng",
            "trong_luong_kg": 3,
        }
    ]
    resp = await client.post(
        "/api/v1/integration/products", headers={"X-API-Key": key}, json={"rows": rows}
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["created"] == 1
    page = (await client.get("/api/v1/products", headers=auth(admin_token))).json()
    assert page["items"][0]["code"] == "API-1"

    bad = await client.post(
        "/api/v1/integration/products", headers={"X-API-Key": key + "x"}, json={"rows": rows}
    )
    assert bad.status_code == 401

    key_id = resp_key_id = (await client.get("/api/v1/api-keys", headers=auth(admin_token))).json()[
        0
    ]["id"]
    assert resp_key_id
    await client.post(f"/api/v1/api-keys/{key_id}/revoke", headers=auth(admin_token))
    revoked = await client.post(
        "/api/v1/integration/products", headers={"X-API-Key": key}, json={"rows": rows}
    )
    assert revoked.status_code == 401


def test_split_parts_value_allocation() -> None:
    class P:
        def __init__(self, temp: str, weight: str) -> None:
            from app.modules.master_data.models import TempClass

            self.temp_class = TempClass(temp)
            self.weight_kg = Decimal(weight)
            self.id = None

    lines = [
        LineDraft(P("DRY", "1"), Decimal(1), None),  # type: ignore[arg-type]
        LineDraft(P("FROZEN", "2"), Decimal(1), None),  # type: ignore[arg-type]
    ]
    parts = split_parts(lines, Decimal("100"))
    assert parts[TempGroup.DRY] == (Decimal(1), Decimal("33.33"))
    assert parts[TempGroup.COLD] == (Decimal(2), Decimal("66.67"))
