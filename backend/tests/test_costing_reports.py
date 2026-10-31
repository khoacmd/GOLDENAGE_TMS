from decimal import Decimal
from io import BytesIO
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook

from app.modules.costing.service import split
from tests.conftest import PLAN_DATE, Env, auth, stop, trip_body
from tests.test_delivery import dispatched_trip, item_ids


def test_split_rounds_to_whole_dong() -> None:
    assert split(Decimal(100), [Decimal(1)] * 3) == [Decimal(34), Decimal(33), Decimal(33)]
    assert split(Decimal(1000), [Decimal(0), Decimal(0)]) == [Decimal(500), Decimal(500)]
    parts = split(Decimal(305_001), [Decimal("12.5"), Decimal("80"), Decimal("7.333")])
    assert sum(parts) == 305_001 and parts[1] > parts[0] > parts[2]
    assert split(Decimal(10), []) == []


async def finish(client: AsyncClient, env: Env, trip: dict[str, Any], fail_first: bool) -> None:
    for n, item in enumerate(item_ids(trip)):
        failed = fail_first and n == 0
        body = {"status": "FAILED", "reason": "Khách vắng nhà"} if failed else {"status": "SUCCESS"}
        resp = await client.patch(
            f"/api/v1/trips/{trip['id']}/items/{item}", headers=auth(env.manager), json=body
        )
        assert resp.status_code == 200, resp.text


async def closed_trip(client: AsyncClient, env: Env) -> dict[str, Any]:
    trip = await dispatched_trip(client, env)
    await finish(client, env, trip, fail_first=True)
    resp = await client.post(f"/api/v1/trips/{trip['id']}/close", headers=auth(env.manager))
    assert resp.json()["order_status"] == "CLOSED", resp.text
    return trip


async def test_close_records_cost_and_allocates_by_weight(client: AsyncClient, env: Env) -> None:
    trip = await closed_trip(client, env)
    cost = (await client.get(f"/api/v1/trips/{trip['id']}/cost", headers=auth(env.manager))).json()
    assert cost["distance_source"] == "ESTIMATE" and float(cost["distance_km"]) > 0
    assert cost["allocation_method"] == "WEIGHT"  # XE02 trong dữ liệu mẫu
    allocations = cost["allocations"]
    assert len(allocations) == 3  # cả lượt giao thất bại cũng chịu chi phí
    assert sum(Decimal(a["amount"]) for a in allocations) == Decimal(cost["total_cost"])
    heavy = max(allocations, key=lambda a: Decimal(a["weight_kg"]))
    assert heavy["amount"] == max(a["amount"] for a in allocations)


async def test_manager_records_odometer_and_extra_cost(client: AsyncClient, env: Env) -> None:
    trip = await closed_trip(client, env)
    url = f"/api/v1/trips/{trip['id']}/cost"
    body = {"odometer_start": 1000, "odometer_end": 1042.5, "extra_cost": 50000, "note": "Phí cầu"}
    cost = (await client.put(url, headers=auth(env.manager), json=body)).json()
    assert cost["distance_source"] == "ODOMETER"
    assert Decimal(cost["distance_km"]) == Decimal("42.5")
    assert Decimal(cost["total_cost"]) == Decimal("42.5") * 6000 + 50000
    assert sum(Decimal(a["amount"]) for a in cost["allocations"]) == Decimal(cost["total_cost"])

    bad = {"odometer_start": 1050, "odometer_end": 1000}
    resp = await client.put(url, headers=auth(env.manager), json=bad)
    assert resp.status_code == 400 and resp.json()["error"]["code"] == "invalid_odometer"
    resp = await client.put(url, headers=auth(env.keeper), json=body)
    assert resp.status_code == 403


async def test_per_trip_method_splits_equally_per_order(client: AsyncClient, env: Env) -> None:
    await client.patch(
        f"/api/v1/vehicles/{env.vehicles['XE02']}",
        headers=auth(env.admin),
        json={"cost_allocation_method": "PER_TRIP"},
    )
    trip = await closed_trip(client, env)
    cost = (
        await client.put(
            f"/api/v1/trips/{trip['id']}/cost", headers=auth(env.manager), json={"distance_km": 20}
        )
    ).json()
    per_order: dict[str, Decimal] = {}
    for a in cost["allocations"]:
        per_order[a["external_ref"]] = per_order.get(a["external_ref"], Decimal(0)) + Decimal(
            a["amount"]
        )
    assert cost["allocation_method"] == "PER_TRIP" and len(per_order) == 2
    assert abs(per_order["DH2610-001"] - per_order["DH2610-007"]) <= 1


async def test_cost_only_after_dispatch(client: AsyncClient, env: Env) -> None:
    body = trip_body(env, "XE02", "KHO01", [stop(env, ("DH2610-007", "DRY"))])
    trip = (await client.post("/api/v1/trips", headers=auth(env.manager), json=body)).json()
    resp = await client.put(
        f"/api/v1/trips/{trip['id']}/cost", headers=auth(env.manager), json={"distance_km": 5}
    )
    assert resp.status_code == 400


async def test_report_summary_and_redelivery(client: AsyncClient, env: Env) -> None:
    first = await closed_trip(client, env)
    failed_part = (
        await client.get(f"/api/v1/trips/{first['id']}", headers=auth(env.manager))
    ).json()["stops"][0]["items"][0]
    assert failed_part["status"] == "FAILED"

    # Giao lại phần đơn thất bại vào ngày hôm sau.
    point = env.parts[("DH2610-001", "COLD")]["delivery_point_id"]
    body = trip_body(
        env,
        "XE02",
        "KHO01",
        [{"delivery_point_id": point, "order_part_ids": [failed_part["order_part_id"]]}],
        plan_date="2026-11-03",
    )
    second = (await client.post("/api/v1/trips", headers=auth(env.manager), json=body)).json()
    url = f"/api/v1/trips/{second['id']}"
    await client.post(f"{url}/approve", headers=auth(env.manager))
    await client.post(f"{url}/prepare", headers=auth(env.keeper))
    await client.post(f"{url}/pick", headers=auth(env.keeper))
    second = (await client.post(f"{url}/dispatch", headers=auth(env.driver_token))).json()
    await finish(client, env, second, fail_first=False)
    await client.post(f"{url}/close", headers=auth(env.manager))

    params = {"date_from": PLAN_DATE, "date_to": "2026-11-03"}
    report = (
        await client.get("/api/v1/reports/summary", headers=auth(env.manager), params=params)
    ).json()
    t = report["totals"]
    assert (t["trips"], t["trips_closed"], t["trips_without_cost"]) == (2, 2, 0)
    assert (t["items_success"], t["items_failed"], t["redeliveries"]) == (3, 1, 1)
    assert t["success_rate"] == 0.75
    assert t["orders"] == 2
    assert t["cost"] > 0 and t["cost_per_kg"] == t["cost"] / t["delivered_kg"]
    assert 0 < t["load_utilization"] < 1
    assert [d["plan_date"] for d in report["by_day"]] == [PLAN_DATE, "2026-11-03"]
    assert report["by_vehicle"][0]["vehicle_code"] == "XE02"
    assert report["failure_reasons"] == [{"reason": "Khách vắng nhà", "count": 1}]
    assert sum(c["cost"] for c in report["by_customer"]) == t["cost"]

    resp = await client.get("/api/v1/reports/export", headers=auth(env.manager), params=params)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == [
        "Tổng quan",
        "Theo xe",
        "Theo ngày",
        "Theo khách hàng",
        "Lý do thất bại",
        "Chi tiết đơn",
    ]
    assert wb["Chi tiết đơn"].max_row == 1 + 4

    assert (
        await client.get("/api/v1/reports/summary", headers=auth(env.driver_token), params=params)
    ).status_code == 403
    bad = {"date_from": "2026-11-05", "date_to": "2026-11-01"}
    resp = await client.get("/api/v1/reports/summary", headers=auth(env.manager), params=bad)
    assert resp.status_code == 400
